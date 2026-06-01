"""Extract a high-level scope list + property metadata from a PIP file.

Lighter-weight than the full PIP-to-Budget extractor in `services/extractor.py`:
the scheduler only needs the high-level scope buckets (Guestrooms, Lobby,
F&B, …) to seed the Scope Selection page, plus property name / brand / keys
for the Property Context page. Line-item granularity isn't useful at this layer.
"""
from __future__ import annotations

import json
import logging
from pathlib import Path
from typing import Any

from services.claude_client import complete_json, pdf_block, text_block, ClaudeError
from services.extractor import _extract_docx_text


logger = logging.getLogger(__name__)


# Mirrors STANDARD_SCOPE in frontend ScheduleGenerator.tsx — keep these
# strings byte-identical so the boolean map merges cleanly into wizard state.
SCOPE_KEYS = [
    "Guestrooms",
    "Guest Corridors",
    "Lobby & Public Areas",
    "F&B Outlets",
    "Meeting & Function Space",
    "Spa & Wellness",
    "Pool & Outdoor Amenities",
    "Fitness Center",
    "Retail",
    "Back of House",
    "MEP / Building Systems",
    "Building Envelope / Facade",
    "Vertical Transportation",
    "Structural / Recert Work",
    "Site / Landscape",
]


_SYSTEM = f"""You are reading a hotel PIP (Property Improvement Plan) or similar scope document. Extract:
1. Property metadata (best effort — null if not stated).
2. The high-level scope buckets that are addressed by this PIP.

Return ONLY a JSON object — no preamble, no markdown fences:
{{
  "property": {{
    "name": string or null,
    "brand": string or null,
    "location": string or null,         // "City, State"
    "keys": integer or null,
    "propertyType": one of ["Urban", "Resort", "Branded Residential"] or null
  }},
  "scope": {{
    {", ".join(f'"{k}": boolean' for k in SCOPE_KEYS)}
  }},
  "notes": "<one short paragraph (≤ 80 words) summarizing the scope, special requirements, and any key dates the PIP mentions>"
}}

Bucket-mapping rules:
- Guestrooms ↔ any guestroom upgrade (hard goods, soft goods, case goods, bathrooms, AV).
- Guest Corridors ↔ corridor finishes / lighting / signage.
- Lobby & Public Areas ↔ arrival, lobby, lounge, club lounge, guest elevator landings.
- F&B Outlets ↔ restaurants, bars, lounges, room service.
- Meeting & Function Space ↔ ballrooms, breakout rooms, prefunction.
- Spa & Wellness ↔ spa treatment / locker / hydrotherapy.
- Pool & Outdoor Amenities ↔ pools, decks, cabanas, outdoor F&B.
- Fitness Center ↔ fitness, gym, group class space.
- Retail ↔ retail, sundries, gift shop.
- Back of House ↔ employee dining, locker rooms, kitchens (BOH portion), admin.
- MEP / Building Systems ↔ chillers, boilers, AHUs, electrical, plumbing, fire alarm, BMS.
- Building Envelope / Facade ↔ facade, roof, windows, waterproofing.
- Vertical Transportation ↔ elevators, escalators (modernization or replacement).
- Structural / Recert Work ↔ recert (25/40/50 yr), GPR-driven repairs, balcony/garage structural.
- Site / Landscape ↔ porte cochere, drives, landscaping, signage at site level.

Set a bucket to true ONLY if there's evidence the document calls for work in that area.
If the PIP is silent on a bucket, return false (not null)."""


def extract_pip_for_schedule(file_path: Path) -> dict[str, Any]:
    """Read the file, send it to Claude, return the parsed scope + property
    payload normalized to the frontend's expected shape."""
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        block = pdf_block(file_path)
    elif suffix == ".docx":
        block = text_block(_extract_docx_text(file_path))
    elif suffix in (".xlsx", ".xls"):
        block = text_block(_xlsx_to_text(file_path))
    elif suffix in (".txt", ".md"):
        block = text_block(file_path.read_text(encoding="utf-8", errors="replace"))
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    raw = complete_json(
        system=_SYSTEM,
        user_content=[block],
        max_tokens=2000,
    )

    # Normalize — guarantee every scope key exists so the frontend can spread
    # the result over the existing wizard state without missing entries.
    scope = {k: bool((raw.get("scope") or {}).get(k, False)) for k in SCOPE_KEYS}
    prop = raw.get("property") or {}
    out = {
        "property": {
            "name": prop.get("name"),
            "brand": prop.get("brand"),
            "location": prop.get("location"),
            "keys": prop.get("keys"),
            "propertyType": prop.get("propertyType"),
        },
        "scope": scope,
        "notes": str(raw.get("notes") or ""),
    }
    return out


def _xlsx_to_text(path: Path) -> str:
    """Flatten an xlsx into a plain-text representation. Useful for owner's
    matrix / scope spreadsheets that Claude can read just as easily as text."""
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"### Sheet: {sheet_name}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                parts.append(" | ".join(cells))
        parts.append("")
    return "\n".join(parts)[:60_000]  # cap to keep token cost predictable
