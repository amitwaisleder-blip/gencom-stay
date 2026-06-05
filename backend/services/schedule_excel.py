"""Excel writer for Schedule Generator output.

Matches the layout of `RCCP - Draft ID Schedule - 04.13.26.xlsx`:
  Sheet 1 — ID Schedule:
    Row 1: project name banner
    Row 3: header row (TASK / MILESTONE | TYPE | OVERRIDE | START | END | WKS | <weekly columns>)
    Row 4: weekly day numbers under their month headers
    Body: phase headers (caps, bold, gold band) interleaved with task rows.
          Task rows include a shaded Gantt bar across the weekly grid.
  Sheet 2 — Milestones:
    Simple 4-col list (Milestone | Start | End | Duration)
"""
from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from services.schedule_engine import Schedule, Task


GENCOM_GREEN = "FF2D5A27"
GENCOM_GOLD  = "FFB89555"
LIGHT_SAND   = "FFF5F3EE"
HEADER_GRAY  = "FFE8E8E8"
PHASE_FILL   = "FFEFE7D2"  # warm gold tint for phase headers
BAR_FILL     = "FF6FA88A"  # green Gantt bar
INDEPENDENT_BAR = "FFB89555"  # gold for Independent tasks
MILESTONE_FILL = "FFB89555"

THIN = Side(style="thin", color="FFD9D4C8")
BORDER_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BORDER_BOTTOM = Border(bottom=Side(style="medium", color=GENCOM_GREEN))


# Fixed-info columns (1-indexed):
COL_TASK     = 1
COL_TYPE     = 2
COL_OVERRIDE = 3
COL_START    = 4
COL_END      = 5
COL_WKS      = 6
GRID_FIRST   = 7  # first weekly column


def _monday(d: date) -> date:
    return d - timedelta(days=d.weekday())


def _fmt(d: date) -> str:
    """US-style MM/DD/YYYY for spreadsheet readability — ISO is unfriendly to
    non-technical readers and many of the schedule's consumers (owners, GCs)
    expect month-first."""
    return d.strftime("%m/%d/%Y")


def _build_week_grid(start: date, end: date) -> list[date]:
    """Return a list of Monday dates spanning [start, end] inclusive — used as
    the column anchors for the Gantt grid."""
    cur = _monday(start)
    last = _monday(end) + timedelta(days=7)
    out: list[date] = []
    while cur <= last:
        out.append(cur)
        cur += timedelta(days=7)
    return out


def write_schedule_xlsx(sched: Schedule) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    ws.title = "ID Schedule"

    # Determine grid bounds from the task list.
    all_starts = [t.start for t in sched.tasks]
    all_ends   = [t.end   for t in sched.tasks]
    if not all_starts:
        # Empty schedule fallback — write headers only.
        all_starts = [date.today()]
        all_ends   = [date.today() + timedelta(days=7)]
    grid = _build_week_grid(min(all_starts), max(all_ends))

    # ─── banner row ──────────────────────────────────────────────────────
    ws.cell(1, 1, sched.project_name).font = Font(name="Arial", size=14, bold=True, color=GENCOM_GREEN)
    ws.cell(1, 1).alignment = Alignment(horizontal="left", vertical="center")

    # ─── header row 3 (column names) ────────────────────────────────────
    headers = [
        (COL_TASK, "TASK / MILESTONE"),
        (COL_TYPE, "TYPE"),
        (COL_OVERRIDE, "OVERRIDE"),
        (COL_START, "START"),
        (COL_END, "END"),
        (COL_WKS, "WKS"),
    ]
    for col, label in headers:
        c = ws.cell(3, col, label)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=GENCOM_GREEN)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_ALL

    # ─── month + day header rows (rows 2 and 4) ─────────────────────────
    # Row 2: month name spanning each block of weeks within a month
    # Row 4: day-of-month numbers for each weekly column
    prev_month = None
    span_start = None
    for i, monday in enumerate(grid):
        col = GRID_FIRST + i
        # Day-of-month tag in row 4
        day_cell = ws.cell(4, col, monday.day)
        day_cell.font = Font(name="Arial", size=8, color=GENCOM_GREEN)
        day_cell.alignment = Alignment(horizontal="center")
        day_cell.fill = PatternFill("solid", fgColor=LIGHT_SAND)
        # Month banner in row 2 — merge across same-month columns
        if monday.month != prev_month:
            if span_start is not None and span_start < col:
                ws.merge_cells(start_row=2, start_column=span_start, end_row=2, end_column=col - 1)
                m_cell = ws.cell(2, span_start)
                m_cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
                m_cell.fill = PatternFill("solid", fgColor=GENCOM_GOLD)
                m_cell.alignment = Alignment(horizontal="center", vertical="center")
            ws.cell(2, col, monday.strftime("%b %Y"))
            span_start = col
            prev_month = monday.month
    # close the trailing month span
    if span_start is not None and span_start < GRID_FIRST + len(grid):
        ws.merge_cells(start_row=2, start_column=span_start, end_row=2,
                       end_column=GRID_FIRST + len(grid) - 1)
        m_cell = ws.cell(2, span_start)
        m_cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
        m_cell.fill = PatternFill("solid", fgColor=GENCOM_GOLD)
        m_cell.alignment = Alignment(horizontal="center", vertical="center")

    # ─── body rows ──────────────────────────────────────────────────────
    row = 5
    for t in sched.tasks:
        _write_task_row(ws, row, t, grid)
        row += 1

    # ─── column widths ──────────────────────────────────────────────────
    ws.column_dimensions[get_column_letter(COL_TASK)].width = 38
    ws.column_dimensions[get_column_letter(COL_TYPE)].width = 12
    ws.column_dimensions[get_column_letter(COL_OVERRIDE)].width = 12
    ws.column_dimensions[get_column_letter(COL_START)].width = 12
    ws.column_dimensions[get_column_letter(COL_END)].width = 12
    ws.column_dimensions[get_column_letter(COL_WKS)].width = 7
    for i in range(len(grid)):
        ws.column_dimensions[get_column_letter(GRID_FIRST + i)].width = 3.2

    ws.freeze_panes = ws.cell(5, GRID_FIRST)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 14

    # ─── Milestones sheet ───────────────────────────────────────────────
    ms = wb.create_sheet("Milestones")
    ms.cell(2, 2, "MILESTONE").font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    ms.cell(2, 3, "START").font   = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    ms.cell(2, 4, "END").font     = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    ms.cell(2, 5, "DURATION").font = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    for c in range(2, 6):
        cell = ms.cell(2, c)
        cell.fill = PatternFill("solid", fgColor=GENCOM_GREEN)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER_ALL
    for i, m in enumerate(sched.milestones):
        r = 3 + i
        ms.cell(r, 2, m.name)
        ms.cell(r, 3, _fmt(m.start))
        ms.cell(r, 4, _fmt(m.end))
        ms.cell(r, 5, m.duration)
        for c in range(2, 6):
            ms.cell(r, c).border = BORDER_ALL
            ms.cell(r, c).font = Font(name="Arial", size=10)
    ms.column_dimensions["B"].width = 32
    ms.column_dimensions["C"].width = 14
    ms.column_dimensions["D"].width = 14
    ms.column_dimensions["E"].width = 14
    ms.row_dimensions[2].height = 20

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def _write_task_row(ws, row: int, t: Task, grid: list[date]) -> None:
    is_phase = t.type == "phase_header"
    is_milestone = t.type == "milestone"

    # Task name — indent for sub-rows.
    name_cell = ws.cell(row, COL_TASK, ("    " * t.indent) + t.name)
    name_cell.font = Font(
        name="Arial", size=10,
        bold=is_phase,
        color=GENCOM_GREEN if is_phase else "FF1A1D24",
    )
    name_cell.alignment = Alignment(horizontal="left", vertical="center", indent=t.indent)

    # Type column — only on actual task rows.
    if not is_phase:
        type_label = "Milestone" if is_milestone else t.type.capitalize()
        ws.cell(row, COL_TYPE, type_label).font = Font(name="Arial", size=9, italic=True)
        ws.cell(row, COL_TYPE).alignment = Alignment(horizontal="center")

    # Override / Start / End / Wks
    if t.override_date:
        ws.cell(row, COL_OVERRIDE, _fmt(t.override_date))
    ws.cell(row, COL_START, _fmt(t.start))
    ws.cell(row, COL_END,   _fmt(t.end))
    if is_milestone:
        ws.cell(row, COL_WKS, "—")
    else:
        ws.cell(row, COL_WKS, t.weeks)
    for c in (COL_OVERRIDE, COL_START, COL_END, COL_WKS):
        ws.cell(row, c).font = Font(name="Arial", size=9)
        ws.cell(row, c).alignment = Alignment(horizontal="center")

    # Phase header row gets a tinted band across the info columns.
    if is_phase:
        for c in range(COL_TASK, COL_WKS + 1):
            ws.cell(row, c).fill = PatternFill("solid", fgColor=PHASE_FILL)
            ws.cell(row, c).border = BORDER_BOTTOM

    # ─── Gantt bar shading ──────────────────────────────────────────────
    if not grid:
        return
    bar_color = MILESTONE_FILL if is_milestone else (
        INDEPENDENT_BAR if t.type == "independent" else BAR_FILL
    )
    if is_phase:
        bar_color = GENCOM_GOLD
    bar_fill = PatternFill("solid", fgColor=bar_color)
    for i, monday in enumerate(grid):
        week_end = monday + timedelta(days=6)
        # A week is "in" the task if any overlap exists between [monday, week_end] and [t.start, t.end].
        if t.start <= week_end and t.end >= monday:
            col = GRID_FIRST + i
            ws.cell(row, col).fill = bar_fill
