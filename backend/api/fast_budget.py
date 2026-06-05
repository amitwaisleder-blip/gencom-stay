"""Endpoints supporting the Fast Budget Generator (client-side ROM tool).

Server-side pieces:
  - `enrich` — given a hotel name (plus optional city hint), Claude returns a
    structured payload with brand, location, counts, year built, etc. Used to
    auto-fill the Property Basics step.
  - `scope-description` — given the final budget (Low/Mid/High) and property
    context, Claude returns three high-level scope narratives describing what
    can realistically be purchased / delivered at each budget level.
"""
from __future__ import annotations

import io
import logging
import os
import re
from datetime import date
from typing import Optional, List, Dict

from fastapi import APIRouter, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from services.claude_client import client as claude_client, ClaudeError, _extract_json
from config import settings


logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/fast-budget", tags=["fast-budget"])


class EnrichRequest(BaseModel):
    name: str
    city_hint: Optional[str] = None


class EnrichResponse(BaseModel):
    name: str
    city: Optional[str] = None
    state_or_country: Optional[str] = None
    brand: Optional[str] = None
    tier: Optional[str] = None
    room_count: Optional[int] = None
    suite_count: Optional[int] = None
    floors: Optional[int] = None
    year_built: Optional[int] = None
    last_renovation: Optional[int] = None
    property_type: Optional[str] = None
    gross_sf: Optional[int] = None
    region_key: Optional[str] = None
    confidence: Optional[str] = None  # "high" | "medium" | "low"
    notes: Optional[str] = None


ENRICH_SYSTEM = """You are a hospitality data lookup assistant. Given a hotel name (and optional city hint), return a structured JSON object with known facts about the property. Only use information you are confident about from your training data. For any field you are uncertain about, return null rather than guess.

Return ONLY a JSON object — no preamble, no markdown fences — with these keys:
{
  "name": string (the hotel's canonical name),
  "city": string or null,
  "state_or_country": string or null,   // US state abbreviation (e.g. "CA") or country name if non-US
  "brand": one of [
    "Ritz-Carlton","Four Seasons","Rosewood","St. Regis","Waldorf Astoria","Edition",
    "JW Marriott","W","Westin","Sheraton","Marriott","Hyatt Regency","Thompson",
    "Hyatt Centric","Hyatt Place","InterContinental","Kimpton","Hotel Indigo",
    "Independent/Boutique","Other"
  ] or null,
  "tier": one of ["Luxury","Upper-Upscale","Upscale","Upper-Midscale","Midscale"] or null,
  "room_count": integer or null,
  "suite_count": integer or null,
  "floors": integer or null,
  "year_built": integer or null,
  "last_renovation": integer or null,
  "property_type": one of ["Urban High-Rise","Resort","Airport","Suburban","Conversion"] or null,
  "gross_sf": integer or null,
  "region_key": one of [
    "New York, NY","San Francisco, CA","Honolulu, HI","Miami, FL","Los Angeles, CA",
    "Boston, MA","Washington, DC","Chicago, IL","Seattle, WA","Las Vegas, NV",
    "National Average","Secondary Market","Tertiary Market"
  ] or null,
  "confidence": "high" | "medium" | "low",
  "notes": short string noting any assumptions or ambiguity, or null
}

Rules:
- If two hotels share the name, use the city_hint to disambiguate. If still ambiguous, pick the most prominent one and mention the ambiguity in notes.
- Be conservative. Rough year built or room count is better than a made-up precise number, but null is better than a guess that could be wrong by >20%.
- For region_key, pick the closest match from the allowed list. Major cities map directly. Secondary/tertiary markets map by population tier.
- For unbranded / independent hotels, use "Independent/Boutique".
- Set confidence to "high" only for well-known flagships where most fields are certain."""


@router.post("/enrich", response_model=EnrichResponse)
def enrich(req: EnrichRequest) -> EnrichResponse:
    if not req.name.strip():
        raise HTTPException(status_code=400, detail="Hotel name is required")
    if not settings.anthropic_api_key:
        raise HTTPException(
            status_code=503,
            detail="Claude API key not configured. Add ANTHROPIC_API_KEY in backend/.env.",
        )

    user_text = f"Hotel name: {req.name.strip()}"
    if req.city_hint and req.city_hint.strip():
        user_text += f"\nCity hint: {req.city_hint.strip()}"

    try:
        c = claude_client()
        with c.messages.stream(
            model=settings.extraction_model,
            max_tokens=1500,
            system=ENRICH_SYSTEM,
            messages=[{"role": "user", "content": user_text}],
        ) as stream:
            final = stream.get_final_message()
        text = "".join(
            b.text for b in final.content if getattr(b, "type", None) == "text"
        )
        data = _extract_json(text)
    except ClaudeError as e:
        raise HTTPException(status_code=502, detail=f"Claude error: {e}")
    except Exception as e:
        logger.exception("Fast-budget enrich failed")
        raise HTTPException(status_code=500, detail=str(e))

    # Normalize — pydantic drops unknown keys anyway, but coerce obvious int
    # fields (Claude occasionally returns strings).
    for k in ("room_count", "suite_count", "floors", "year_built", "last_renovation", "gross_sf"):
        v = data.get(k)
        if isinstance(v, str):
            try:
                data[k] = int(v.replace(",", "").strip())
            except ValueError:
                data[k] = None

    return EnrichResponse(**data)


# ---------------------------------------------------------------------------
# Scope description — Claude narrates what the Low / Mid / High budget buys.
# ---------------------------------------------------------------------------


class BudgetLevel(BaseModel):
    total: float
    per_key: float
    hard: float
    soft: float


class ScopeDescriptionRequest(BaseModel):
    name: str
    city: Optional[str] = None
    state_or_country: Optional[str] = None
    brand: Optional[str] = None
    tier: Optional[str] = None
    property_type: Optional[str] = None
    room_count: Optional[int] = None
    suite_count: Optional[int] = None
    year_built: Optional[int] = None
    last_renovation: Optional[int] = None
    included_areas: List[str] = []  # human-readable area names
    budgets: Dict[str, BudgetLevel]  # keys: "low" | "mid" | "high"


class ScopeDescriptionResponse(BaseModel):
    low: str
    mid: str
    high: str


SCOPE_DESCRIPTION_SYSTEM = """You are a hospitality renovation consultant writing high-level scope narratives for an acquisitions committee. Given a hotel and three budget levels (Low / Mid / High), write three separate scope descriptions — one per budget — describing in general terms what work CAN typically be delivered at each level for a property of this size, tier, and property type.

Ground your narrative in facts you are given: year built, last renovated, key count, brand, tier, property type, city, and the list of scope areas the user has opted into. Use those facts to calibrate how much scope a budget can plausibly cover.

STRICT RULES — read carefully, these matter:

1. DO NOT invent facts about the hotel that were not provided. You have not been given a PCA, a Property Condition Assessment, an engineering report, owner notes, or any on-the-ground inspection findings. You do not know the actual condition of the roof, elevators, MEP, or any specific system. Never say or imply you are working from a PCA, inspection, assessment, report, walkthrough, or any document other than the scope areas and budget the user gave you.

2. DO NOT commit to specific deferred-maintenance scope as if it is required. Do not say "the roof will be replaced", "elevators will be modernized", "the chiller must be replaced", or similar definite statements. Scope at each budget level is illustrative of what this budget typically supports — not a commitment that any particular system needs work.

3. Use conditional / illustrative phrasing: "typically supports", "can cover", "allows for", "is generally sufficient for", "may include", "usually scoped at this level". Never "will", "must", "required", "confirmed".

4. Stay general and high-level. Describe the TYPE of work the budget enables (e.g. "selective soft-goods refresh across guestrooms and public areas" instead of naming specific systems). Avoid specifics that would require knowing condition.

5. Unless owner notes or an inspection are explicitly referenced in the input, do not single out individual systems as needing work. It is fine to say "MEP scope at this level typically allows for targeted equipment replacement" but NOT "the HVAC is near end of life and needs replacement".

6. Lead each description with what the budget level supports in concrete categories (finishes, soft goods, FF&E, selective vs. full system work). Note what is typically OUT of scope at each level.

7. Do NOT quote dollar figures — the cost card next to your text shows those.

8. 2–4 short paragraphs OR a short paragraph + 4–8 bullets, under ~200 words per level. Plain active language, no marketing fluff.

Return ONLY a JSON object — no preamble, no markdown fences:
{
  "low": "...",
  "mid": "...",
  "high": "..."
}
"""


def _fmt_budget(label: str, lvl: BudgetLevel) -> str:
    return (
        f"  - {label}: total ${lvl.total:,.0f} "
        f"(${lvl.per_key:,.0f}/key · hard ${lvl.hard:,.0f} · soft ${lvl.soft:,.0f})"
    )


@router.post("/scope-description", response_model=ScopeDescriptionResponse)
def scope_description(req: ScopeDescriptionRequest) -> ScopeDescriptionResponse:
    if not req.name.strip():
        raise HTTPException(status_code=400, detail="Property name is required")
    if "low" not in req.budgets or "mid" not in req.budgets or "high" not in req.budgets:
        raise HTTPException(
            status_code=400,
            detail="budgets must include low, mid, and high entries",
        )
    if not settings.anthropic_api_key:
        raise HTTPException(
            status_code=503,
            detail="Claude API key not configured. Add ANTHROPIC_API_KEY in backend/.env.",
        )

    # Compose a compact user message with all the facts Claude needs.
    address = ", ".join(filter(None, [req.city, req.state_or_country])) or "—"
    header_lines = [
        f"Property: {req.name}",
        f"Location: {address}",
        f"Brand / Tier: {req.brand or '—'} / {req.tier or '—'}",
        f"Type: {req.property_type or '—'}",
        f"Keys: {req.room_count if req.room_count is not None else '—'}"
        + (f" (suites: {req.suite_count})" if req.suite_count else ""),
        f"Year built: {req.year_built if req.year_built is not None else '—'}",
        f"Last renovated: {req.last_renovation if req.last_renovation is not None else '—'}",
    ]
    budget_lines = [
        "Budgets:",
        _fmt_budget("Low", req.budgets["low"]),
        _fmt_budget("Mid", req.budgets["mid"]),
        _fmt_budget("High", req.budgets["high"]),
    ]
    area_line = (
        "Included scope areas: " + ", ".join(req.included_areas)
        if req.included_areas
        else "Included scope areas: (none specified — assume typical full-service mix)"
    )
    user_text = "\n".join(header_lines + [area_line] + budget_lines)

    try:
        c = claude_client()
        # Stream to sidestep the 10-minute guard on larger calls.
        with c.messages.stream(
            model=settings.extraction_model,
            max_tokens=4000,
            system=SCOPE_DESCRIPTION_SYSTEM,
            messages=[{"role": "user", "content": user_text}],
        ) as stream:
            final = stream.get_final_message()
        text = "".join(
            b.text for b in final.content if getattr(b, "type", None) == "text"
        )
        data = _extract_json(text)
    except ClaudeError as e:
        raise HTTPException(status_code=502, detail=f"Claude error: {e}")
    except Exception as e:
        logger.exception("Fast-budget scope-description failed")
        raise HTTPException(status_code=500, detail=str(e))

    for k in ("low", "mid", "high"):
        if not isinstance(data.get(k), str):
            data[k] = ""

    return ScopeDescriptionResponse(low=data["low"], mid=data["mid"], high=data["high"])


# ---------------------------------------------------------------------------
# PPTX export — generates a basic slide deck from the final ROM figures.
# When a branded template ships, drop it at
# `backend/templates/fast_budget_template.pptx` and the endpoint will use it
# instead of building from scratch. Placeholder slide matching is still done
# by search-and-replace on text tokens like `{{PROPERTY_NAME}}`.
# ---------------------------------------------------------------------------

TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
    "templates",
    "fast_budget_template.pptx",
)


class PptxLine(BaseModel):
    area: str
    scope: str
    low: float
    mid: float
    high: float
    category: str  # "dm" | "interior" | "custom"


class PptxSoftRow(BaseModel):
    label: str
    pct: Optional[float] = None
    basis: Optional[str] = None
    low: float
    mid: float
    high: float


class PptxSection(BaseModel):
    text: str  # scope narrative for a single budget level (optional)


class PptxExportRequest(BaseModel):
    name: str
    city: Optional[str] = None
    state_or_country: Optional[str] = None
    brand: Optional[str] = None
    tier: Optional[str] = None
    property_type: Optional[str] = None
    room_count: Optional[int] = None
    suite_count: Optional[int] = None
    year_built: Optional[int] = None
    last_renovation: Optional[int] = None
    region_key: Optional[str] = None
    hard_range: Dict[str, float]  # low/mid/high
    softs_range: Dict[str, float]  # excl. dev fee
    dev_fee_range: Dict[str, float]
    grand_range: Dict[str, float]
    per_key_range: Dict[str, float]
    dm_subtotal: Dict[str, float]
    interior_subtotal: Dict[str, float]
    lines: List[PptxLine]
    softs: List[PptxSoftRow]
    dev_fee_rows: List[PptxSoftRow]
    scope_low: Optional[str] = None
    scope_mid: Optional[str] = None
    scope_high: Optional[str] = None


def _money(n: float) -> str:
    return f"${n:,.0f}"


def _replace_placeholders(shape, mapping: Dict[str, str]) -> None:
    """Walk a shape's runs and do {{TOKEN}} → value substitutions, preserving
    formatting. Tokens that aren't in the mapping are left alone."""
    if not shape.has_text_frame:
        return
    for para in shape.text_frame.paragraphs:
        for run in para.runs:
            txt = run.text
            if "{{" not in txt:
                continue
            for key, val in mapping.items():
                txt = txt.replace("{{" + key + "}}", val)
            run.text = txt


def _build_token_map(req: PptxExportRequest) -> Dict[str, str]:
    address = ", ".join(filter(None, [req.city, req.state_or_country])) or "—"
    today = date.today().strftime("%m/%d/%y")
    return {
        "PROPERTY_NAME": req.name or "Untitled Property",
        "ADDRESS": address,
        "CITY": req.city or "",
        "STATE_OR_COUNTRY": req.state_or_country or "",
        "BRAND": req.brand or "",
        "TIER": req.tier or "",
        "PROPERTY_TYPE": req.property_type or "",
        "ROOM_COUNT": str(req.room_count) if req.room_count is not None else "",
        "SUITE_COUNT": str(req.suite_count) if req.suite_count is not None else "",
        "YEAR_BUILT": str(req.year_built) if req.year_built is not None else "",
        "LAST_RENOVATION": str(req.last_renovation) if req.last_renovation is not None else "",
        "REGION_KEY": req.region_key or "",
        "DATE": today,
        "GRAND_LOW": _money(req.grand_range["low"]),
        "GRAND_MID": _money(req.grand_range["mid"]),
        "GRAND_HIGH": _money(req.grand_range["high"]),
        "PER_KEY_LOW": _money(req.per_key_range["low"]),
        "PER_KEY_MID": _money(req.per_key_range["mid"]),
        "PER_KEY_HIGH": _money(req.per_key_range["high"]),
        "HARD_LOW": _money(req.hard_range["low"]),
        "HARD_MID": _money(req.hard_range["mid"]),
        "HARD_HIGH": _money(req.hard_range["high"]),
        "SOFT_LOW": _money(req.softs_range["low"]),
        "SOFT_MID": _money(req.softs_range["mid"]),
        "SOFT_HIGH": _money(req.softs_range["high"]),
        "DEV_FEE_LOW": _money(req.dev_fee_range["low"]),
        "DEV_FEE_MID": _money(req.dev_fee_range["mid"]),
        "DEV_FEE_HIGH": _money(req.dev_fee_range["high"]),
        "DM_SUB_LOW": _money(req.dm_subtotal["low"]),
        "DM_SUB_MID": _money(req.dm_subtotal["mid"]),
        "DM_SUB_HIGH": _money(req.dm_subtotal["high"]),
        "INT_SUB_LOW": _money(req.interior_subtotal["low"]),
        "INT_SUB_MID": _money(req.interior_subtotal["mid"]),
        "INT_SUB_HIGH": _money(req.interior_subtotal["high"]),
        "SCOPE_LOW": req.scope_low or "",
        "SCOPE_MID": req.scope_mid or "",
        "SCOPE_HIGH": req.scope_high or "",
    }


def _build_from_template(req: PptxExportRequest, template_path: str) -> io.BytesIO:
    from pptx import Presentation  # local import keeps module load cheap

    prs = Presentation(template_path)
    mapping = _build_token_map(req)
    for slide in prs.slides:
        for shape in slide.shapes:
            _replace_placeholders(shape, mapping)
    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def _build_from_scratch(req: PptxExportRequest) -> io.BytesIO:
    """Fallback PPTX — used until the branded template is dropped in place.
    Creates a minimal cover + summary + details deck so the button works today."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.dml.color import RGBColor

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    BLANK = prs.slide_layouts[6]
    INK = RGBColor(0x1A, 0x1D, 0x24)
    STONE = RGBColor(0x6B, 0x6F, 0x78)
    SAND = RGBColor(0xD9, 0xD4, 0xC8)
    GOLD = RGBColor(0xB8, 0x95, 0x55)

    def add_text(slide, left, top, width, height, text, *, size=14, bold=False,
                 color=INK, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = align
        r = p.add_run()
        r.text = text
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
        return box

    # ---- Slide 1: cover ----------------------------------------------------
    s = prs.slides.add_slide(BLANK)
    add_text(s, Inches(0.5), Inches(0.4), Inches(2), Inches(0.4),
             "GENCOM", size=14, bold=True, color=GOLD)
    add_text(s, Inches(0.5), Inches(2.8), Inches(12.3), Inches(0.9),
             req.name or "Untitled Property", size=40, bold=True,
             align=PP_ALIGN.CENTER)
    add_text(s, Inches(0.5), Inches(3.8), Inches(12.3), Inches(0.5),
             "ROM CAPEX Budget", size=16, color=STONE, align=PP_ALIGN.CENTER)
    address = ", ".join(filter(None, [req.city, req.state_or_country])) or ""
    site = " · ".join(filter(None, [
        address,
        f"{req.room_count} keys" if req.room_count else "",
        req.brand or "",
        req.tier or "",
    ]))
    if site:
        add_text(s, Inches(0.5), Inches(4.4), Inches(12.3), Inches(0.5),
                 site, size=12, color=STONE, align=PP_ALIGN.CENTER)
    add_text(s, Inches(11.5), Inches(0.4), Inches(1.3), Inches(0.4),
             date.today().strftime("%m/%d/%y"), size=11, color=STONE,
             align=PP_ALIGN.RIGHT)

    # ---- Slide 2: Estimate summary ----------------------------------------
    s = prs.slides.add_slide(BLANK)
    add_text(s, Inches(0.5), Inches(0.4), Inches(8), Inches(0.4),
             "Estimate Summary", size=18, bold=True)
    card_w = Inches(4.0)
    card_h = Inches(3.6)
    card_y = Inches(1.5)
    labels = [
        ("Low Estimate", req.grand_range["low"], req.hard_range["low"], req.softs_range["low"], req.dev_fee_range["low"], req.per_key_range["low"]),
        ("Mid Estimate", req.grand_range["mid"], req.hard_range["mid"], req.softs_range["mid"], req.dev_fee_range["mid"], req.per_key_range["mid"]),
        ("High Estimate", req.grand_range["high"], req.hard_range["high"], req.softs_range["high"], req.dev_fee_range["high"], req.per_key_range["high"]),
    ]
    for i, (label, total, hard, soft, dev, pk) in enumerate(labels):
        left = Inches(0.5 + i * 4.25)
        # Card border
        shp = s.shapes.add_shape(1, left, card_y, card_w, card_h)  # rectangle
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF) if label != "Mid Estimate" else RGBColor(0xEC, 0xFD, 0xF5)
        shp.line.color.rgb = SAND
        # Text
        add_text(s, left + Inches(0.25), card_y + Inches(0.2), card_w - Inches(0.5), Inches(0.4),
                 label.upper(), size=10, bold=True, color=STONE)
        add_text(s, left + Inches(0.25), card_y + Inches(0.6), card_w - Inches(0.5), Inches(0.9),
                 _money(total), size=28, bold=True)
        add_text(s, left + Inches(0.25), card_y + Inches(1.7), card_w - Inches(0.5), Inches(0.4),
                 f"Hard costs  {_money(hard)}", size=11, color=STONE)
        add_text(s, left + Inches(0.25), card_y + Inches(2.1), card_w - Inches(0.5), Inches(0.4),
                 f"Soft costs  {_money(soft)}", size=11, color=STONE)
        add_text(s, left + Inches(0.25), card_y + Inches(2.5), card_w - Inches(0.5), Inches(0.4),
                 f"Dev fee     {_money(dev)}", size=11, color=STONE)
        add_text(s, left + Inches(0.25), card_y + Inches(3.0), card_w - Inches(0.5), Inches(0.4),
                 f"Per key  {_money(pk)}", size=12, bold=True)

    # ---- Slide 3: Hard Cost breakdown table -------------------------------
    s = prs.slides.add_slide(BLANK)
    add_text(s, Inches(0.5), Inches(0.4), Inches(10), Inches(0.4),
             "Hard Costs by Area", size=18, bold=True)
    rows = [("Area", "Scope", "Low", "Mid", "High")]
    for l in req.lines:
        rows.append((l.area, l.scope, _money(l.low), _money(l.mid), _money(l.high)))
    rows.append(("Hard Cost Subtotal", "", _money(req.hard_range["low"]),
                 _money(req.hard_range["mid"]), _money(req.hard_range["high"])))
    table_shape = s.shapes.add_table(
        rows=len(rows), cols=5,
        left=Inches(0.5), top=Inches(1.1),
        width=Inches(12.3), height=Inches(0.3 * len(rows)),
    )
    tbl = table_shape.table
    for c_i, head in enumerate(rows[0]):
        cell = tbl.cell(0, c_i)
        cell.text = head
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = INK
    for r_i, row in enumerate(rows[1:], start=1):
        is_subtotal = r_i == len(rows) - 1
        for c_i, val in enumerate(row):
            cell = tbl.cell(r_i, c_i)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.bold = is_subtotal
                    run.font.color.rgb = INK

    # ---- Slide 4: Scope narratives (if present) ---------------------------
    if req.scope_low or req.scope_mid or req.scope_high:
        s = prs.slides.add_slide(BLANK)
        add_text(s, Inches(0.5), Inches(0.4), Inches(10), Inches(0.4),
                 "High-Level Scope Description", size=18, bold=True)
        col_w = Inches(4.1)
        y = Inches(1.1)
        for i, (label, text, total) in enumerate([
            ("Low", req.scope_low or "", req.grand_range["low"]),
            ("Mid", req.scope_mid or "", req.grand_range["mid"]),
            ("High", req.scope_high or "", req.grand_range["high"]),
        ]):
            left = Inches(0.5 + i * 4.25)
            add_text(s, left, y, col_w, Inches(0.4),
                     f"{label}  {_money(total)}", size=14, bold=True)
            add_text(s, left, y + Inches(0.5), col_w, Inches(5.5),
                     text, size=10, color=INK)

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


@router.post("/export-pptx")
def export_pptx(req: PptxExportRequest):
    try:
        if os.path.exists(TEMPLATE_PATH):
            buf = _build_from_template(req, TEMPLATE_PATH)
            source = "template"
        else:
            buf = _build_from_scratch(req)
            source = "generated"
    except Exception as e:
        logger.exception("Fast-budget PPTX export failed")
        raise HTTPException(status_code=500, detail=f"PPTX generation failed: {e}")

    safe = re.sub(r"[^\w\-]+", "_", req.name or "property").strip("_") or "property"
    today = date.today().strftime("%Y-%m-%d")
    filename = f"{safe}_ROM_CAPEX_{today}.pptx"
    logger.info("Fast-budget PPTX export (%s) → %s", source, filename)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
